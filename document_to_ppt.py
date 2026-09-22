import requests
import json
import sys
from pptx import Presentation

OLLAMA_URL = "http://127.0.0.1:11434/api/generate"
MODEL = "promotex:latest"


# Get document text from command-line argument
document_text = sys.argv[1]

prompt = f"""
Create a professional PowerPoint presentation based ONLY on this document.

DOCUMENT:
-------------------------
{document_text}
-------------------------

Return ONLY valid JSON in this format:

{{
  "title": "Presentation title",
  "slides": [
    {{
      "title": "Slide title",
      "points": [
        "Point 1",
        "Point 2",
        "Point 3"
      ]
    }}
  ]
}}

Create exactly 5 slides.
Each slide should have 3 to 5 useful points.
Do not invent facts.
Do not add markdown or explanations.
"""


# Ask Promotex AI
response = requests.post(
    OLLAMA_URL,
    json={
        "model": MODEL,
        "prompt": prompt,
        "stream": False
    }
)

result = response.json()

# Convert AI response into Python data
presentation_data = json.loads(result["response"])


# Create PowerPoint
prs = Presentation()

# Title slide
slide = prs.slides.add_slide(prs.slide_layouts[0])

slide.shapes.title.text = presentation_data["title"]
slide.placeholders[1].text = "Generated locally by Promotex AI"


# Content slides
for slide_data in presentation_data["slides"]:

    slide = prs.slides.add_slide(
        prs.slide_layouts[1]
    )

    slide.shapes.title.text = slide_data["title"]

    body = slide.placeholders[1].text_frame

    body.text = slide_data["points"][0]

    for point in slide_data["points"][1:]:

        paragraph = body.add_paragraph()
        paragraph.text = point


# Save presentation
output_file = "promotex_document_generated.pptx"

prs.save(output_file)

print("\nPPT created successfully!")
print(f"File: {output_file}")