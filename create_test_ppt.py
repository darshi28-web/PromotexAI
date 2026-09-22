from pptx import Presentation

prs = Presentation()

# Title slide
slide = prs.slides.add_slide(prs.slide_layouts[0])

slide.shapes.title.text = "Promotex AI"
slide.placeholders[1].text = "Local AI Workbench"

# Second slide
slide = prs.slides.add_slide(prs.slide_layouts[1])

slide.shapes.title.text = "Key Capabilities"

body = slide.placeholders[1].text_frame

body.text = "General AI Chat"

for item in [
    "Private document analysis",
    "Accreditation gap analysis",
    "Local AI processing",
    "Offline-ready architecture",
]:
    p = body.add_paragraph()
    p.text = item

# Third slide
slide = prs.slides.add_slide(prs.slide_layouts[1])

slide.shapes.title.text = "Education Use Case"

body = slide.placeholders[1].text_frame

body.text = "Confidential Academic Document Management"

for item in [
    "Upload accreditation documents",
    "Analyze evidence",
    "Identify missing evidence",
    "Generate reports and presentations",
]:
    p = body.add_paragraph()
    p.text = item

# Save
prs.save("promotex_test.pptx")

print("PPT created successfully!")
print("File: promotex_test.pptx")